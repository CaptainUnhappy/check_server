"""Create a Word report from inspection results."""

from datetime import datetime, timezone
from pathlib import Path
import zipfile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def build_word(results, output, failed_hosts=()):
    doc = Document()
    # python-docx's bundled empty document includes a bibliography customXml part.
    # Remove it and the thumbnail; neither is needed for this report.
    for rel_id, rel in list(doc.part.rels.items()):
        if rel.reltype.endswith("/customXml"):
            doc.part.drop_rel(rel_id)
    for rel_id, rel in list(doc.part.package.rels.items()):
        if rel.reltype.endswith("/thumbnail"):
            del doc.part.package.rels[rel_id]
    for border in doc.styles.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.65)
    section.left_margin = section.right_margin = Inches(0.7)
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.color.rgb = RGBColor(0, 0, 0)
    doc.styles["Normal"].font.size = Pt(10)
    doc.styles["Normal"].paragraph_format.space_after = Pt(5)
    doc.styles["Title"].font.size = Pt(20)
    props = doc.core_properties
    props.author = props.last_modified_by = "check_server"
    props.title = "服务器巡检示例报告"
    props.subject = "示例输出，需人工复核"
    props.comments = "Generated from HTML inspection reports."
    props.created = props.modified = datetime.now(timezone.utc)
    for index, result in enumerate(results):
        if index:
            doc.add_page_break()
        doc.add_heading("服务器巡检示例报告", 0)
        doc.add_paragraph("本报告汇总资源阈值检查结果，并列出待人工复核事项。通过仅表示该项指标满足规则。")
        doc.add_paragraph(f"主机标识  {result['host_alias']}    采集时间  {result['collected_at'] or '未知'}")
        doc.add_paragraph("数据性质  " + ("合成演示数据" if result["synthetic"] else
                                     "示例数据" if result.get("example") else "授权采集数据 仅限内部使用"))
        if failed_hosts:
            doc.add_paragraph("采集不完整 以下主机失败，不能据此判断整批正常：" + ", ".join(failed_hosts))
        table = doc.add_table(rows=1, cols=4)
        table.autofit = False
        widths = (2.1, 0.8, 0.8, 3.4)
        for col, width in zip(table.columns, widths):
            col.width = Inches(width)
        for cell, label in zip(table.rows[0].cells, ("检查项", "值", "结果", "规则或后续动作")):
            cell.text = label
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "DCE6F1")
            cell._tc.get_or_add_tcPr().append(shade)
        repeat = OxmlElement("w:tblHeader")
        table.rows[0]._tr.get_or_add_trPr().append(repeat)
        for item in result["checks"]:
            cells = table.add_row().cells
            value_text = "—" if item["value"] is None else f"{item['value']:g}"
            if item["value"] is not None and item["item"] in ("CPU使用率", "内存使用率", "磁盘最高占用率"):
                value_text += "%"
            for cell, value in zip(cells, (item["item"], value_text, item["status"], item["rule"])):
                cell.text = value
        for row in table.rows:
            for index, (cell, width) in enumerate(zip(row.cells, widths)):
                cell.width = Inches(width)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcpr = cell._tc.get_or_add_tcPr()
                borders = OxmlElement("w:tcBorders")
                for edge in ("top", "left", "bottom", "right"):
                    border = OxmlElement("w:" + edge)
                    for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                        border.set(qn("w:" + key), value)
                    borders.append(border)
                tcpr.append(borders)
                for paragraph in cell.paragraphs:
                    if index in (1, 2):
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    paragraph.paragraph_format.space_before = Pt(3)
                    paragraph.paragraph_format.space_after = Pt(3)
        doc.add_paragraph("复核人  ______________    复核日期  ______________")
        doc.add_paragraph("处理意见  __________________________________________________________")
        doc.add_paragraph(f"附件见交付包 html/{result['host_alias']}.html，与本报告一并交付。")
    doc.save(output)


def write_bundle(output, html_paths, results):
    output = Path(output)
    with zipfile.ZipFile(output / "delivery.zip", "w", zipfile.ZIP_DEFLATED) as bundle:
        for name in ("inspection.docx", "summary.json"):
            bundle.write(output / name, name)
        for path, result in zip(html_paths, results):
            bundle.write(path, f"html/{result['host_alias']}.html")
